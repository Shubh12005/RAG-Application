import os
from pathlib import Path
import docx
import pandas as pd
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

sample_dir = Path(__file__).resolve().parent.parent / "sample_documents"
sample_dir.mkdir(parents=True, exist_ok=True)

def generate_pdf():
    pdf_path = sample_dir / "Research_Paper.pdf"
    doc = SimpleDocTemplate(str(pdf_path), pagesize=letter)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=15
    )
    heading_style = ParagraphStyle(
        'DocHeading',
        parent=styles['Heading2'],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#334155'),
        spaceBefore=12,
        spaceAfter=8
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceAfter=10
    )

    story = []

    # Page 1: Introduction & Main Objective
    story.append(Paragraph("Scalable Retrieval-Augmented Generation for macOS Environments", title_style))
    story.append(Paragraph("<b>Authors:</b> Advanced AI Systems Laboratory", body_style))
    story.append(Spacer(1, 15))
    story.append(Paragraph("1. Executive Summary & Project Objectives", heading_style))
    story.append(Paragraph(
        "The primary objective of this project is to develop an autonomous, low-latency "
        "Retrieval-Augmented Generation (RAG) system tailored for macOS workstations. "
        "Modern knowledge workers require instantaneous semantic querying across diverse local documents "
        "without leaking proprietary corporate intelligence to public servers. The system architecture "
        "prioritizes strict anti-hallucination guardrails, verifiable page-level citation provenance, "
        "and local vector storage.",
        body_style
    ))
    story.append(Paragraph(
        "Key performance indicators include sub-200ms semantic similarity search across millions of embedding "
        "dimensions, accurate page preservation across complex PDF layouts, and an intuitive research workspace UI.",
        body_style
    ))
    story.append(PageBreak())

    # Page 2: System Architecture
    story.append(Paragraph("2. System Architecture & Component Design", heading_style))
    story.append(Paragraph(
        "The system architecture integrates a decoupled client-server paradigm. The backend is orchestrated "
        "using Python and FastAPI, executing asynchronous non-blocking event loops for file ingestion. "
        "Text extraction leverages PyPDF for PDF document hierarchies, python-docx for Office Open XML files, "
        "and Pandas for structured tabular CSV ingestion.",
        body_style
    ))
    story.append(Paragraph(
        "Embeddings are computed locally using SentenceTransformers (specifically the all-MiniLM-L6-v2 model), "
        "which produces 384-dimensional dense vectors. These vectors are indexed into Chroma, an open-source "
        "embedded vector database that persists HNSW index files locally on disk.",
        body_style
    ))
    story.append(PageBreak())

    # Page 3: Evaluation & Expected Results
    story.append(Paragraph("3. Empirical Evaluation and Expected Results", heading_style))
    story.append(Paragraph(
        "Empirical benchmarks demonstrate that combining top-k semantic retrieval (k=5) with recursive "
        "character text splitting (chunk size 800 characters, overlap 100 characters) achieves a 98.4% retrieval "
        "precision on standard research benchmarks. Compared to naive full-document prompting, the RAG approach "
        "reduces token consumption by 87% while completely eliminating hallucinated citations.",
        body_style
    ))
    story.append(Paragraph(
        "In failure mode analyses where queried facts are absent from the ingested corpus, the system's "
        "anti-hallucination system prompt reliably triggers a standardized refusal response: "
        "'I could not find this information in the uploaded document.'",
        body_style
    ))
    story.append(PageBreak())

    # Page 4: Conclusions
    story.append(Paragraph("4. Conclusions and Future Research", heading_style))
    story.append(Paragraph(
        "We have demonstrated a production-ready RAG application that operates seamlessly on Apple Silicon "
        "hardware. The local persistence engine guarantees state preservation across restarts, allowing researchers "
        "to maintain long-term indexed corpora with zero external cloud dependencies.",
        body_style
    ))

    doc.build(story)
    print(f"Generated PDF at: {pdf_path}")

def generate_docx():
    docx_path = sample_dir / "Project_Specifications.docx"
    doc = docx.Document()
    doc.add_heading("Hardware Acceleration & System Specifications", 0)

    doc.add_heading("1. Target Hardware Environment", level=1)
    doc.add_paragraph(
        "The application is engineered specifically for macOS systems running on Apple Silicon (M1, M2, M3, M4) "
        "and Intel architectures. PyTorch's Metal Performance Shaders (MPS) backend is dynamically probed "
        "upon initialization to accelerate vector matrix multiplications."
    )

    doc.add_heading("2. Storage and Data Privacy", level=1)
    doc.add_paragraph(
        "All uploaded documents are stored in the local file directory backend/data/uploads/, and persistent "
        "vector indexes are maintained in backend/data/vectorstore/. No user documents or embedding vectors "
        "are ever transmitted to external vector hosting providers."
    )

    doc.add_heading("3. Supported File Formats", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Format"
    hdr_cells[1].text = "Parser Library"
    hdr_cells[2].text = "Page Tracking"

    rows = [
        ("PDF (.pdf)", "PyPDF", "Yes (1-based physical pages)"),
        ("Word (.docx)", "python-docx", "Yes (Logical section & heading pages)"),
        ("CSV (.csv)", "Pandas", "Yes (Batch row groupings)"),
        ("Markdown (.md)", "Custom Markdown Parser", "Yes (Header sections)"),
    ]
    for fmt, lib, pt in rows:
        row_cells = table.add_row().cells
        row_cells[0].text = fmt
        row_cells[1].text = lib
        row_cells[2].text = pt

    doc.save(str(docx_path))
    print(f"Generated DOCX at: {docx_path}")

def generate_csv():
    csv_path = sample_dir / "benchmarks_dataset.csv"
    data = {
        "Model": ["all-MiniLM-L6-v2", "bge-small-en-v1.5", "all-mpnet-base-v2", "text-embedding-3-small"],
        "Dimension": [384, 384, 768, 1536],
        "Latency_ms": [18.2, 22.4, 45.1, 85.0],
        "Memory_MB": [85, 120, 420, 0],
        "Local_Execution": ["Yes", "Yes", "Yes", "No (API)"],
        "Accuracy_Score": [0.942, 0.951, 0.968, 0.958]
    }
    df = pd.DataFrame(data)
    df.to_csv(csv_path, index=False)
    print(f"Generated CSV at: {csv_path}")

def generate_markdown():
    md_path = sample_dir / "anti_hallucination_guidelines.md"
    content = """# Guidelines for Anti-Hallucination in RAG Systems

## 1. Core Principle
Retrieval-Augmented Generation systems must never invent facts. When a user asks a question, the LLM must synthesize the answer using **only** the retrieved context passages provided in the prompt.

## 2. Refusal Protocol
If the retrieved chunks do not contain sufficient evidence to answer the query, the system must trigger the canonical refusal sentence:
> "I could not find this information in the uploaded document."

## 3. Provenance and Citation
Every statement made in the answer should correspond to a specific chunk and page reference:
- PDF files: cite the document title and 1-based page number.
- DOCX and Markdown: cite the heading or section name.
- CSV: cite the row numbers and column names.
"""
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Generated Markdown at: {md_path}")

if __name__ == "__main__":
    generate_pdf()
    generate_docx()
    generate_csv()
    generate_markdown()
    print("All sample test documents generated successfully!")
